"""Gemeinsame Fixtures: erzeugt harmlose Beispieldateien im Temp-Ordner."""
import sys
from pathlib import Path

import pytest

ROOT = Path(__file__).resolve().parent.parent
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))


@pytest.fixture
def jpeg_file(tmp_path):
    from PIL import Image
    p = tmp_path / "foto.jpg"
    img = Image.new("RGB", (64, 48), (120, 30, 200))
    for x in range(64):
        img.putpixel((x, x % 48), (255, 255, 0))
    img.save(p, "JPEG", quality=85)
    return p


@pytest.fixture
def pdf_file(tmp_path):
    import pikepdf
    p = tmp_path / "doc.pdf"
    pdf = pikepdf.new()
    pdf.add_blank_page(page_size=(200, 200))
    pdf.save(p)
    return p


@pytest.fixture
def docx_file(tmp_path):
    import docx
    p = tmp_path / "brief.docx"
    d = docx.Document()
    d.add_paragraph("Hallo Welt")
    d.save(p)
    return p


@pytest.fixture
def xlsx_file(tmp_path):
    import openpyxl
    p = tmp_path / "tabelle.xlsx"
    wb = openpyxl.Workbook()
    wb.active["A1"] = "x"
    wb.save(p)
    return p


@pytest.fixture(autouse=True)
def _isolated_signing_key(tmp_path, monkeypatch):
    """Tests nie mit dem echten Arbeitsplatz-Schlüssel in data/signing arbeiten lassen."""
    try:
        import core.signing as signing
    except Exception:
        return
    monkeypatch.setattr(signing, "KEY_DIR", tmp_path / "signing_keys")

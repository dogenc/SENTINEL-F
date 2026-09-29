"""Dokument-Stammbaum: PDF-/ID, XMP-Abstammung, Word-RSIDs."""
import os
import re
import shutil
import zipfile

import pikepdf
import pytest

from core.history import HistoryDB
from core.lineage import relatives
from engines import run_forensic
from engines.analyzers.lineage import xmp_lineage, pdf_lineage


def _docx(tmp_path, name, rsids, root="00A1B2C3"):
    import docx
    p = tmp_path / name
    docx.Document().save(p)
    tmp = tmp_path / (name + ".tmp")
    with zipfile.ZipFile(p) as zin, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for zi in zin.infolist():
            data = zin.read(zi)
            if zi.filename == "word/settings.xml":
                block = f'<w:rsids><w:rsidRoot w:val="{root}"/>' + "".join(
                    f'<w:rsid w:val="{r}"/>' for r in rsids) + "</w:rsids>"
                data = re.sub(rb"<w:rsids>.*?</w:rsids>", block.encode(), data, flags=re.S)
                if b"<w:rsids>" not in data:
                    data = data.replace(b"</w:settings>", block.encode() + b"</w:settings>")
            zout.writestr(zi, data)
    shutil.move(tmp, p)
    return p


def _pdf(path, ident, saves=1):
    pdf = pikepdf.new()
    pdf.add_blank_page()
    pdf.save(path, static_id=True)
    data = path.read_bytes()
    data = re.sub(rb"/ID \[ ?<[0-9a-f]+> ?<[0-9a-f]+> ?\]", b"/ID [<" + ident + b"> <" + os.urandom(8).hex().encode() + b">]", data)
    data += b"\n%%EOF\n" * (saves - 1)
    path.write_bytes(data)
    return path


def test_extractors():
    assert pdf_lineage(b"trailer << /ID [<AABBCCDD11223344> <99887766>] >>\n%%EOF") == \
        {"pdf_id0": "aabbccdd11223344", "pdf_id1": "99887766", "pdf_saves": 1}
    x = xmp_lineage(b'<x:xmpmeta xmpMM:DocumentID="xmp.did:ABC12345" xmpMM:InstanceID="xmp.iid:INST0001" '
                    b'xmpMM:OriginalDocumentID="xmp.did:ORIG0001"><stRef:documentID>xmp.did:PARENT01</stRef:documentID>'
                    b'<stEvt:instanceID>xmp.iid:OLD00001</stEvt:instanceID></x:xmpmeta>')
    assert x == {"xmp_doc": "abc12345", "xmp_orig": "orig0001", "xmp_inst": "inst0001",
                 "xmp_derived": ["parent01"], "xmp_hist": ["old00001"]}


def test_word_versions_and_branch(tmp_path):
    db = HistoryDB(tmp_path / "h.db")
    v1 = db.record(run_forensic(str(_docx(tmp_path, "vertrag_v1.docx", ["00A1B2C3", "00B00001"]))))
    v2 = db.record(run_forensic(str(_docx(tmp_path, "vertrag_v2.docx", ["00A1B2C3", "00B00001", "00C00002"]))))
    br = db.record(run_forensic(str(_docx(tmp_path, "vertrag_anwalt.docx", ["00A1B2C3", "00B00001", "00D00003"]))))
    other = db.record(run_forensic(str(_docx(tmp_path, "fremd.docx", ["11111111"], root="11111111"))))
    rel = {r["name"]: r["relation"] for r in relatives(db, v2)}
    assert rel == {"vertrag_v1.docx": "ancestor", "vertrag_anwalt.docx": "sibling"}
    assert {r["name"]: r["relation"] for r in relatives(db, v1)} == \
        {"vertrag_v2.docx": "descendant", "vertrag_anwalt.docx": "descendant"}
    assert relatives(db, other) == []
    assert "rsidRoot 00A1B2C3" in relatives(db, br)[0]["evidence"]


def test_pdf_versions(tmp_path):
    db = HistoryDB(tmp_path / "h.db")
    a = db.record(run_forensic(str(_pdf(tmp_path / "angebot.pdf", b"cafebabe00112233"))))
    b = db.record(run_forensic(str(_pdf(tmp_path / "angebot_geaendert.pdf", b"cafebabe00112233", saves=2))))
    db.record(run_forensic(str(_pdf(tmp_path / "anderes.pdf", b"0000111122223333"))))
    assert [(r["name"], r["relation"]) for r in relatives(db, b)] == [("angebot.pdf", "ancestor")]
    assert [(r["name"], r["relation"]) for r in relatives(db, a)] == [("angebot_geaendert.pdf", "descendant")]


def test_xmp_derived_image(tmp_path):
    from PIL import Image
    import struct

    def jpg(name, xmp):
        p = tmp_path / name
        Image.new("RGB", (32, 32)).save(p, "JPEG")
        d = p.read_bytes()
        payload = b"http://ns.adobe.com/xap/1.0/\x00" + xmp
        p.write_bytes(d[:2] + b"\xff\xe1" + struct.pack(">H", len(payload) + 2) + payload + d[2:])
        return p
    db = HistoryDB(tmp_path / "h.db")
    src = db.record(run_forensic(str(jpg("original.jpg", b'<x xmpMM:DocumentID="xmp.did:SRC00001" '
                                                         b'xmpMM:InstanceID="xmp.iid:SRC00001"/>'))))
    edit = db.record(run_forensic(str(jpg("retusche.jpg", b'<x xmpMM:DocumentID="xmp.did:EDT00002" '
                                                          b'xmpMM:InstanceID="xmp.iid:EDT00002">'
                                                          b'<stRef:documentID>xmp.did:SRC00001</stRef:documentID></x>'))))
    assert [(r["name"], r["relation"]) for r in relatives(db, edit)] == [("original.jpg", "ancestor")]
    assert [(r["name"], r["relation"]) for r in relatives(db, src)] == [("retusche.jpg", "descendant")]


def test_lineage_tab(tmp_path):
    os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
    QtWidgets = pytest.importorskip("PyQt6.QtWidgets")
    app = QtWidgets.QApplication.instance() or QtWidgets.QApplication([])  # noqa: F841 – Referenz halten
    from gui.analysis_view import AnalysisView
    db = HistoryDB(tmp_path / "h.db")
    db.record(run_forensic(str(_docx(tmp_path, "v1.docx", ["00A1B2C3"]))))
    res = run_forensic(str(_docx(tmp_path, "v2.docx", ["00A1B2C3", "00C00002"])))
    v2 = db.record(res)
    view = AnalysisView()
    view.load_result(res)
    view.set_lineage(relatives(db, v2))
    assert "LINEAGE  1" in view.tabs.tabText(view.tabs.indexOf(view.lineage))
    assert view.lineage.tree.topLevelItem(0).text(0).startswith("▲")
    assert "rsid_root" in view.lineage.lbl.text()


def test_template_twins_are_not_related(tmp_path):
    """Zwei unabhängige Dokumente aus derselben Vorlage (identische RSIDs) sind keine Verwandten."""
    import docx
    db = HistoryDB(tmp_path / "h.db")
    ids = []
    for n in ("brief_a.docx", "brief_b.docx"):
        d = docx.Document()
        d.add_paragraph(f"Inhalt {n}")
        d.save(tmp_path / n)
        ids.append(db.record(run_forensic(str(tmp_path / n))))
    assert relatives(db, ids[0]) == []


def test_generic_ids_are_ignored(tmp_path, monkeypatch):
    import core.lineage as cl
    monkeypatch.setattr(cl, "GENERIC_LIMIT", 2)
    db = HistoryDB(tmp_path / "h.db")
    ids = [db.record(run_forensic(str(_pdf(tmp_path / f"f{i}.pdf", b"deadbeef00000000", saves=i + 1))))
           for i in range(4)]
    assert relatives(db, ids[0]) == []

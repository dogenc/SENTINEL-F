"""
Sandbox-Worker: läuft als eigener Prozess mit Low Integrity im Job Object (core/sandbox.py).

Nur hier werden Dateiinhalte geparst. Ergebnisse gehen ausschließlich als Dateien
in den eigenen Job-Ordner (%LOCALAPPDATA%Low):
    progress.jsonl   Fortschritt für die GUI
    result.json      vollständiges Analyseergebnis
    preview.png      NEU kodierte, verkleinerte Vorschau (Bild / 1. PDF-Seite)
    preview.txt      Textvorschau (Office/PDF)
    worker.log       stdout/stderr des Workers

Keine Qt-Imports: der Worker braucht keine Oberfläche.
"""
import json
import sys
import time
import warnings
from pathlib import Path

PREVIEW_MAX = 1600


def _quiet_errors():
    """Keine Windows-Fehlerdialoge aus dem Worker (Absturz = Befund, nicht Popup)."""
    if sys.platform == "win32":
        import ctypes
        ctypes.windll.kernel32.SetErrorMode(0x0001 | 0x0002 | 0x8000)


def _harden_parsers():
    from PIL import Image
    # Bild-Bomben: Pixel-Obergrenze (Pillow-Standard) wird zum harten Fehler
    Image.MAX_IMAGE_PIXELS = 89_478_485
    warnings.simplefilter("error", Image.DecompressionBombWarning)


def _progress_writer(jobdir):
    f = open(jobdir / "progress.jsonl", "a", encoding="utf-8", buffering=1)

    def progress(label, value):
        f.write(json.dumps({"label": str(label)[:120], "value": int(value), "t": time.time()}) + "\n")
    return progress


def _preview(path, result, jobdir):
    kind = result.get("kind")
    ext = Path(path).suffix.lower()
    try:
        if kind == "image":
            from PIL import Image
            with Image.open(path) as im:
                im.thumbnail((PREVIEW_MAX, PREVIEW_MAX))
                im.convert("RGB").save(jobdir / "preview.png", "PNG")
        elif kind == "pdf":
            import fitz
            with fitz.open(path) as doc:
                if doc.page_count:
                    page = doc[0]
                    zoom = min(1.5, PREVIEW_MAX / max(1.0, page.rect.width, page.rect.height))
                    page.get_pixmap(matrix=fitz.Matrix(zoom, zoom), alpha=False).save(str(jobdir / "preview.png"))
                    text = page.get_text()[:15000]
                    (jobdir / "preview.txt").write_text(f"Seiten: {doc.page_count}\n\n{text}", encoding="utf-8")
        elif kind == "document":
            (jobdir / "preview.txt").write_text(_document_text(path, ext), encoding="utf-8")
    except Exception as e:   # Vorschau ist optional – Analyse bleibt gültig
        (jobdir / "preview.txt").write_text(f"(Vorschau nicht möglich: {type(e).__name__}: {e})", encoding="utf-8")


def _document_text(path, ext):
    lines = []
    if ext in (".docx", ".docm", ".dotx"):
        import docx
        d = docx.Document(path)
        lines += [p.text for p in d.paragraphs[:400] if p.text.strip()]
    elif ext in (".xlsx", ".xlsm"):
        import openpyxl
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        for ws in wb.worksheets[:6]:
            lines.append(f"── {ws.title} ──")
            for i, row in enumerate(ws.iter_rows(values_only=True)):
                lines.append(" | ".join("" if c is None else str(c) for c in row[:10])[:400])
                if i >= 40:
                    break
    elif ext in (".pptx", ".pptm"):
        from pptx import Presentation
        pres = Presentation(path)
        for i, slide in enumerate(list(pres.slides)[:40], 1):
            lines.append(f"── Folie {i} ──")
            for shape in slide.shapes:
                if shape.has_text_frame:
                    lines += [p.text for p in shape.text_frame.paragraphs if p.text.strip()]
    else:
        lines.append(f"Textvorschau für {ext} nicht verfügbar – siehe Analyse.")
    return "\n".join(lines)[:20000]


def analyze(path, jobdir):
    from engines import run_forensic
    from engines.analyzers.base import json_safe
    progress = _progress_writer(jobdir)
    res = run_forensic(path, progress_fn=progress)
    progress("Sichere Vorschau", 98)
    if res.get("ok"):
        _preview(path, res, jobdir)
    tmp = jobdir / "result.json.part"
    tmp.write_text(json.dumps(json_safe(res), ensure_ascii=False), encoding="utf-8")
    tmp.replace(jobdir / "result.json")
    return 0


# ─── Selbsttest-Sonden (beweisen, dass die Sandbox blockiert) ───────────────
def selftest(probe, jobdir):
    out = {"probe": probe}
    try:
        if probe == "spawn":
            import subprocess
            subprocess.run(["cmd.exe", "/c", "echo", "ausgebrochen"], capture_output=True, timeout=10)
            out.update(blocked=False, detail="Kindprozess wurde gestartet")
        elif probe == "write":
            target = Path.home() / f"sentinel_probe_{jobdir.name}.txt"
            target.write_text("probe")
            target.unlink()
            out.update(blocked=False, detail=f"Schreiben nach {target.parent} möglich")
        elif probe == "registry":
            import winreg
            k = winreg.CreateKey(winreg.HKEY_CURRENT_USER, r"Software\DGKN_Sentinel_Probe")
            winreg.CloseKey(k)
            winreg.DeleteKey(winreg.HKEY_CURRENT_USER, r"Software\DGKN_Sentinel_Probe")
            out.update(blocked=False, detail="Registry-Schreiben möglich")
        elif probe == "memory":
            chunks = [bytearray(256 << 20) for _ in range(16)]   # 4 GiB
            out.update(blocked=False, detail=f"{len(chunks) * 256} MiB belegt")
        elif probe == "network":
            import socket
            socket.create_connection(("1.1.1.1", 443), timeout=5).close()
            out.update(blocked=False, detail="TCP-Verbindung ins Internet möglich")
        elif probe == "read_user":
            docs = [p for p in (Path.home() / "Documents").glob("*") if p.is_file()][:1]
            if not docs:
                out.update(blocked=None, detail="keine Datei in Documents zum Testen")
            else:
                docs[0].read_bytes()[:1]
                out.update(blocked=False, detail=f"Nutzerdatei lesbar: {docs[0].name}")
        elif probe == "sleep":
            time.sleep(600)
            out.update(blocked=False, detail="nicht beendet")
        else:
            out.update(blocked=None, detail="unbekannte Sonde")
    except (OSError, MemoryError, PermissionError) as e:
        out.update(blocked=True, detail=f"{type(e).__name__}: {e}")
    (jobdir / "selftest.json").write_text(json.dumps(out), encoding="utf-8")
    return 0


def diff(a, b, jobdir):
    from engines.diff import content_diff
    from engines.analyzers.base import json_safe
    res = content_diff(a, b, out_png=str(jobdir / "diff.png"))
    tmp = jobdir / "diff.json.part"
    tmp.write_text(json.dumps(json_safe(res), ensure_ascii=False), encoding="utf-8")
    tmp.replace(jobdir / "diff.json")
    return 0


def mailbox(path, jobdir):
    from engines.mailbox_split import split
    res = split(path, jobdir / "mails")
    tmp = jobdir / "mailbox.json.part"
    tmp.write_text(json.dumps(res, ensure_ascii=False), encoding="utf-8")
    tmp.replace(jobdir / "mailbox.json")
    return 0


def main(argv):
    from core.sandbox import WORKER_FLAG, SELFTEST_FLAG, DIFF_FLAG, MAILBOX_FLAG
    if argv[0] == DIFF_FLAG:                      # --sentinel-diff A B jobdir
        jobdir = Path(argv[3])
        sys.stdout = sys.stderr = open(jobdir / "worker.log", "a", encoding="utf-8", buffering=1)
        _quiet_errors()
        _harden_parsers()
        return diff(argv[1], argv[2], jobdir)
    mode, arg, jobdir = argv[0], argv[1], Path(argv[2])
    if mode == MAILBOX_FLAG:
        sys.stdout = sys.stderr = open(jobdir / "worker.log", "a", encoding="utf-8", buffering=1)
        _quiet_errors()
        return mailbox(arg, jobdir)
    log = open(jobdir / "worker.log", "a", encoding="utf-8", buffering=1)
    sys.stdout = sys.stderr = log
    _quiet_errors()
    if mode == SELFTEST_FLAG:
        return selftest(arg, jobdir)
    if mode == WORKER_FLAG:
        _harden_parsers()
        return analyze(arg, jobdir)
    return 2

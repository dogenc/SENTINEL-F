"""
╔══════════════════════════════════════════════════════════════╗
║          DOCUMENT FORENSIC ENGINE                            ║
║   Extracted & adapted from document_forensics_suite.py       ║
║   Supports: DOCX, XLSX, PPTX, DOC, XLS, PPT, OLE             ║
╚══════════════════════════════════════════════════════════════╝
"""
import os
import hashlib
import datetime
import re
import math
import zipfile, xml.etree.ElementTree as ET
from pathlib import Path

# ─── Dependency Detection ─────────────────────────────────────────────────────
try:
    import PIL  # noqa: F401 – nur Verfügbarkeit
    PIL_OK = True
except ImportError:
    PIL_OK = False

try:
    import docx
    DOCX_OK = True
except ImportError:
    DOCX_OK = False

try:
    import openpyxl
    OPENPYXL_OK = True
except ImportError:
    OPENPYXL_OK = False

try:
    import pptx
    PPTX_OK = True
except ImportError:
    PPTX_OK = False

try:
    import olefile
    OLEFILE_OK = True
except ImportError:
    OLEFILE_OK = False

try:
    from oletools import olevba
    OLETOOLS_OK = True
except ImportError:
    OLETOOLS_OK = False

try:
    import lxml  # noqa: F401 – nur Verfügbarkeit
    LXML_OK = True
except ImportError:
    LXML_OK = False

def human_size(n):
    for unit in ["B", "KB", "MB", "GB"]:
        if n < 1024: return f"{n:.1f} {unit}"
        n /= 1024
    return f"{n:.1f} TB"

def sha256_file(path):
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for chunk in iter(lambda: f.read(8192), b""): h.update(chunk)
    return h.hexdigest()

def md5_file(path):
    h = hashlib.md5()
    with open(path, "rb") as f:
        for chunk in iter(lambda: f.read(8192), b""): h.update(chunk)
    return h.hexdigest()

def entropy(data):
    if not data: return 0.0
    freq = [0] * 256
    for b in data: freq[b] += 1
    ln = len(data)
    ent = 0.0
    for f in freq:
        if f > 0:
            p = f / ln
            ent -= p * math.log2(p)
    return ent

def detect_type(path):
    ext = Path(path).suffix.lower()
    if ext in (".docx", ".docm"): return "docx"
    if ext in (".xlsx", ".xlsm", ".xlsb"): return "xlsx"
    if ext in (".pptx", ".pptm"): return "pptx"
    if ext in (".doc", ".dot"): return "doc"
    if ext in (".xls",): return "xls"
    if ext in (".ppt",): return "ppt"
    # Sniff magic
    with open(path, "rb") as f:
        magic = f.read(8)
    if magic[:4] == b"PK\x03\x04": return "ooxml"
    if magic[:8] == b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1": return "ole"
    return "unknown"


# ═══════════════════════════════════════════════════════════════════════════════
#  DOCUMENT FORENSIC ENGINE
# ═══════════════════════════════════════════════════════════════════════════════

class DocumentForensicEngine:
    NS = {
        "w":   "http://schemas.openxmlformats.org/wordprocessingml/2006/main",
        "r":   "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
        "cp":  "http://schemas.openxmlformats.org/package/2006/metadata/core-properties",
        "dc":  "http://purl.org/dc/elements/1.1/",
        "ep":  "http://schemas.openxmlformats.org/officeDocument/2006/extended-properties",
        "a":   "http://schemas.openxmlformats.org/drawingml/2006/main",
        "p":   "http://schemas.openxmlformats.org/presentationml/2006/main",
        "mc":  "http://schemas.openxmlformats.org/markup-compatibility/2006",
        "dcterms": "http://purl.org/dc/terms/",
        "vt":  "http://schemas.openxmlformats.org/officeDocument/2006/docPropsVTypes",
    }

    def __init__(self, path, log_fn=None):
        self.path = path
        self.log = log_fn or (lambda msg, level="INFO": None)
        self.doc_type = detect_type(path)
        self.findings = []
        self.zip_entries = []
        self.is_ooxml = self.doc_type in ("docx", "xlsx", "pptx", "ooxml")
        self.is_ole = self.doc_type in ("doc", "xls", "ppt", "ole")

    # ── ZIP / OOXML Structure ────────────────────────────────────────────────
    def analyze_zip_structure(self):
        if not self.is_ooxml:
            return []
        entries = []
        try:
            with zipfile.ZipFile(self.path, "r") as z:
                for info in z.infolist():
                    ent = {
                        "name": info.filename,
                        "size": info.file_size,
                        "compressed": info.compress_size,
                        "date": f"{info.date_time[0]}-{info.date_time[1]:02d}-{info.date_time[2]:02d} "
                                f"{info.date_time[3]:02d}:{info.date_time[4]:02d}",
                        "ratio": (1 - info.compress_size / max(info.file_size, 1)) * 100 if info.file_size > 0 else 0,
                        "is_dir": info.filename.endswith("/"),
                    }
                    entries.append(ent)
                    # Flag suspicious entries
                    if info.filename.startswith("..") or "\\" in info.filename:
                        self.findings.append({
                            "type": "ZIP_TRAVERSAL",
                            "detail": f"Suspicious path in archive: {info.filename}",
                            "severity": "high"
                        })
                    if info.file_size > 100_000_000:
                        self.findings.append({
                            "type": "ZIP_BOMB",
                            "detail": f"Very large entry: {info.filename} ({human_size(info.file_size)})",
                            "severity": "medium"
                        })
        except zipfile.BadZipFile:
            self.findings.append({
                "type": "CORRUPT",
                "detail": "File is not a valid ZIP/OOXML archive",
                "severity": "high"
            })
        self.zip_entries = entries
        return entries

    # ── Core / App / Custom XML Metadata ─────────────────────────────────────
    def extract_metadata(self):
        meta = {}
        meta["File"] = os.path.basename(self.path)
        meta["Size"] = human_size(os.path.getsize(self.path))
        meta["Type"] = self.doc_type.upper()
        meta["MD5"] = md5_file(self.path)[:28] + "..."
        meta["SHA256"] = sha256_file(self.path)[:28] + "..."

        if not self.is_ooxml:
            return self._ole_metadata(meta)

        try:
            with zipfile.ZipFile(self.path, "r") as z:
                # core.xml
                for core_path in ["docProps/core.xml", "docprops/core.xml"]:
                    if core_path in z.namelist():
                        xml = z.read(core_path)
                        root = ET.fromstring(xml)
                        mappings = {
                            f"{{{self.NS['dc']}}}title": "Title",
                            f"{{{self.NS['dc']}}}creator": "Creator",
                            f"{{{self.NS['dc']}}}subject": "Subject",
                            f"{{{self.NS['dc']}}}description": "Description",
                            f"{{{self.NS['cp']}}}lastModifiedBy": "Last Modified By",
                            f"{{{self.NS['cp']}}}revision": "Revision",
                            f"{{{self.NS['cp']}}}category": "Category",
                            f"{{{self.NS['dcterms']}}}created": "Created",
                            f"{{{self.NS['dcterms']}}}modified": "Modified",
                            f"{{{self.NS['cp']}}}keywords": "Keywords",
                        }
                        for xpath, label in mappings.items():
                            el = root.find(xpath)
                            if el is not None and el.text:
                                meta[label] = el.text.strip()[:80]
                        break

                # app.xml
                for app_path in ["docProps/app.xml", "docprops/app.xml"]:
                    if app_path in z.namelist():
                        xml = z.read(app_path)
                        root = ET.fromstring(xml)
                        ns = self.NS["ep"]
                        for tag in ["Application", "AppVersion", "Company", "Template",
                                     "TotalTime", "Pages", "Words", "Characters",
                                     "Slides", "Manager", "PresentationFormat"]:
                            el = root.find(f"{{{ns}}}{tag}")
                            if el is not None and el.text:
                                meta[f"App.{tag}"] = el.text.strip()[:80]
                        break

                # custom.xml
                for cust_path in ["docProps/custom.xml", "docprops/custom.xml"]:
                    if cust_path in z.namelist():
                        xml = z.read(cust_path)
                        root = ET.fromstring(xml)
                        for prop in root:
                            name = prop.attrib.get("name", prop.attrib.get("fmtid", "?"))
                            val_el = list(prop)
                            val = val_el[0].text if val_el and val_el[0].text else "—"
                            meta[f"Custom.{name}"] = val[:80]
                        break
        except Exception as e:
            self.log(f"Metadata extraction error: {e}", "WARN")

        # Check metadata consistency
        creator = meta.get("Creator", "")
        last_mod = meta.get("Last Modified By", "")
        if creator and last_mod and creator != last_mod:
            self.findings.append({
                "type": "METADATA",
                "detail": f"Different creator ({creator}) vs last editor ({last_mod})",
                "severity": "low"
            })
        app = meta.get("App.Application", "")
        if app:
            meta["Software"] = app
            # Flag non-Office applications
            suspicious_apps = ["libreoffice", "google", "wps", "onlyoffice"]
            if any(s in app.lower() for s in suspicious_apps):
                self.findings.append({
                    "type": "SOFTWARE",
                    "detail": f"Created with non-Microsoft application: {app}",
                    "severity": "low"
                })

        return meta

    def _ole_metadata(self, meta):
        if not OLEFILE_OK:
            return meta
        try:
            ole = olefile.OleFileIO(self.path)
            ole_meta = ole.get_metadata()
            for attr in ["title", "subject", "author", "keywords", "comments",
                          "last_saved_by", "revision_number", "creating_application",
                          "create_time", "last_saved_time", "num_pages", "num_words"]:
                val = getattr(ole_meta, attr, None)
                if val:
                    if isinstance(val, bytes):
                        val = val.decode("utf-8", "replace")
                    meta[f"OLE.{attr}"] = str(val)[:80]
            ole.close()
        except Exception as e:
            self.log(f"OLE metadata error: {e}", "WARN")
        return meta

    # ── VBA Macro Analysis ───────────────────────────────────────────────────
    def analyze_macros(self):
        results = {
            "has_macros": False,
            "macro_count": 0,
            "suspicious_keywords": [],
            "macro_sources": [],
            "iocs": [],
        }

        # Check extension first
        ext = Path(self.path).suffix.lower()
        if ext in (".docm", ".xlsm", ".pptm", ".doc", ".xls", ".ppt"):
            results["macro_capable"] = True
        else:
            results["macro_capable"] = False

        # OOXML: check for vbaProject.bin inside ZIP
        if self.is_ooxml:
            try:
                with zipfile.ZipFile(self.path, "r") as z:
                    vba_files = [n for n in z.namelist()
                                  if "vbaproject" in n.lower() or "vba" in n.lower()
                                  or n.lower().endswith(".bin") and "macro" in n.lower()]
                    if vba_files:
                        results["has_macros"] = True
                        results["vba_files"] = vba_files
                        self.findings.append({
                            "type": "MACRO",
                            "detail": f"VBA project found: {', '.join(vba_files)}",
                            "severity": "high"
                        })
            except:
                pass

        # oletools deep analysis
        if OLETOOLS_OK:
            try:
                vba_parser = olevba.VBA_Parser(self.path)
                if vba_parser.detect_vba_macros():
                    results["has_macros"] = True
                    for (filename, stream_path, vba_filename, vba_code) in vba_parser.extract_macros():
                        results["macro_count"] += 1
                        results["macro_sources"].append({
                            "file": vba_filename,
                            "stream": stream_path,
                            "code": vba_code,
                            "lines": len(vba_code.splitlines()),
                        })

                    # Analyze for suspicious patterns
                    analysis = vba_parser.analyze_macros()
                    for kw_type, keyword, description in analysis:
                        results["suspicious_keywords"].append({
                            "type": kw_type,
                            "keyword": keyword,
                            "description": description,
                        })
                        if kw_type in ("AutoExec", "Suspicious", "IOC"):
                            severity = "high" if kw_type == "AutoExec" else "medium"
                            self.findings.append({
                                "type": "MACRO_SUSPICIOUS",
                                "detail": f"[{kw_type}] {keyword}: {description}",
                                "severity": severity
                            })
                        if kw_type == "IOC":
                            results["iocs"].append({
                                "type": keyword,
                                "value": description,
                            })
                vba_parser.close()
            except Exception as e:
                self.log(f"oletools VBA analysis: {e}", "WARN")
        elif self.is_ole and OLEFILE_OK:
            # Basic OLE stream check
            try:
                ole = olefile.OleFileIO(self.path)
                for stream in ole.listdir():
                    path_str = "/".join(stream)
                    if "vba" in path_str.lower() or "macro" in path_str.lower():
                        results["has_macros"] = True
                        self.findings.append({
                            "type": "MACRO",
                            "detail": f"VBA stream found: {path_str}",
                            "severity": "high"
                        })
                ole.close()
            except:
                pass

        return results

    # ── Hidden Content Discovery ─────────────────────────────────────────────
    def find_hidden_content(self):
        hidden = {
            "hidden_text": [],
            "hidden_sheets": [],
            "hidden_slides": [],
            "hidden_rows_cols": [],
            "comments": [],
            "tracked_changes": [],
            "white_text": [],
            "tiny_text": [],
            "headers_footers": [],
        }

        if self.doc_type in ("docx", "ooxml") and DOCX_OK:
            self._hidden_docx(hidden)
        elif self.doc_type in ("xlsx",) and OPENPYXL_OK:
            self._hidden_xlsx(hidden)
        elif self.doc_type in ("pptx",) and PPTX_OK:
            self._hidden_pptx(hidden)

        # Also scan raw XML for hidden content
        if self.is_ooxml:
            self._hidden_ooxml_raw(hidden)

        return hidden

    def _hidden_docx(self, hidden):
        try:
            doc = docx.Document(self.path)

            for i, para in enumerate(doc.paragraphs):
                # Hidden text (vanish property)
                for run in para.runs:
                    if run.font and run.font.hidden:
                        hidden["hidden_text"].append({
                            "para": i + 1,
                            "text": run.text[:100],
                            "type": "hidden_font"
                        })
                        self.findings.append({
                            "type": "HIDDEN_TEXT",
                            "detail": f"Hidden text in paragraph {i+1}: {run.text[:60]}",
                            "severity": "medium"
                        })

                    # White text on white background
                    if run.font and run.font.color and run.font.color.rgb:
                        rgb = str(run.font.color.rgb)
                        if rgb.upper() in ("FFFFFF", "FEFEFE", "FDFDFD"):
                            if run.text.strip():
                                hidden["white_text"].append({
                                    "para": i + 1,
                                    "text": run.text[:100],
                                    "color": rgb
                                })
                                self.findings.append({
                                    "type": "WHITE_TEXT",
                                    "detail": f"White/invisible text: {run.text[:60]}",
                                    "severity": "high"
                                })

                    # Very small text
                    if run.font and run.font.size:
                        pt = run.font.size.pt if hasattr(run.font.size, 'pt') else 0
                        if 0 < pt < 4 and run.text.strip():
                            hidden["tiny_text"].append({
                                "para": i + 1,
                                "text": run.text[:100],
                                "size_pt": pt
                            })
                            self.findings.append({
                                "type": "TINY_TEXT",
                                "detail": f"Microscopic text ({pt}pt): {run.text[:60]}",
                                "severity": "medium"
                            })

            # Comments
            if hasattr(doc, 'part') and doc.part:
                try:
                    comments_part = doc.part.package.part_related_by(
                        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/comments"
                    )
                    if comments_part:
                        xml = comments_part.blob
                        root = ET.fromstring(xml)
                        for comment in root.findall(f".//{{{self.NS['w']}}}comment"):
                            author = comment.attrib.get(f"{{{self.NS['w']}}}author", "?")
                            date = comment.attrib.get(f"{{{self.NS['w']}}}date", "?")
                            texts = []
                            for t in comment.findall(f".//{{{self.NS['w']}}}t"):
                                if t.text: texts.append(t.text)
                            if texts:
                                hidden["comments"].append({
                                    "author": author,
                                    "date": date,
                                    "text": " ".join(texts)[:200]
                                })
                except:
                    pass

            # Headers / Footers
            for section in doc.sections:
                for hf_type, hf_name in [(section.header, "Header"),
                                           (section.footer, "Footer")]:
                    if hf_type and hf_type.is_linked_to_previous is False:
                        for para in hf_type.paragraphs:
                            if para.text.strip():
                                hidden["headers_footers"].append({
                                    "type": hf_name,
                                    "text": para.text[:100]
                                })
        except Exception as e:
            self.log(f"DOCX hidden content scan: {e}", "WARN")

    def _hidden_xlsx(self, hidden):
        try:
            wb = openpyxl.load_workbook(self.path, data_only=True)

            for sheet_name in wb.sheetnames:
                ws = wb[sheet_name]

                # Hidden sheets
                if ws.sheet_state != "visible":
                    hidden["hidden_sheets"].append({
                        "name": sheet_name,
                        "state": ws.sheet_state,
                    })
                    self.findings.append({
                        "type": "HIDDEN_SHEET",
                        "detail": f"Hidden sheet: '{sheet_name}' (state: {ws.sheet_state})",
                        "severity": "high"
                    })

                # Hidden rows/columns
                hidden_rows = []
                hidden_cols = []
                if ws.row_dimensions:
                    for row_idx, rd in ws.row_dimensions.items():
                        if rd.hidden:
                            hidden_rows.append(row_idx)
                if ws.column_dimensions:
                    for col_letter, cd in ws.column_dimensions.items():
                        if cd.hidden:
                            hidden_cols.append(col_letter)

                if hidden_rows:
                    hidden["hidden_rows_cols"].append({
                        "sheet": sheet_name,
                        "hidden_rows": hidden_rows[:20],
                        "total_hidden_rows": len(hidden_rows),
                    })
                    self.findings.append({
                        "type": "HIDDEN_ROWS",
                        "detail": f"Sheet '{sheet_name}': {len(hidden_rows)} hidden row(s)",
                        "severity": "medium"
                    })

                if hidden_cols:
                    hidden["hidden_rows_cols"].append({
                        "sheet": sheet_name,
                        "hidden_cols": hidden_cols[:20],
                        "total_hidden_cols": len(hidden_cols),
                    })

                # Comments
                for row in ws.iter_rows():
                    for cell in row:
                        if cell.comment:
                            hidden["comments"].append({
                                "sheet": sheet_name,
                                "cell": cell.coordinate,
                                "author": cell.comment.author or "?",
                                "text": str(cell.comment.text)[:200],
                            })

                        # White text
                        if cell.font and cell.font.color:
                            try:
                                rgb = cell.font.color.rgb
                                if rgb and str(rgb).upper() in ("00FFFFFF", "FFFFFF", "FFFFFFFF"):
                                    if cell.value:
                                        hidden["white_text"].append({
                                            "sheet": sheet_name,
                                            "cell": cell.coordinate,
                                            "value": str(cell.value)[:100],
                                        })
                                        self.findings.append({
                                            "type": "WHITE_TEXT",
                                            "detail": f"White text at {sheet_name}!{cell.coordinate}: {str(cell.value)[:60]}",
                                            "severity": "high"
                                        })
                            except:
                                pass
            wb.close()
        except Exception as e:
            self.log(f"XLSX hidden content scan: {e}", "WARN")

    def _hidden_pptx(self, hidden):
        try:
            prs = pptx.Presentation(self.path)
            for i, slide in enumerate(prs.slides):
                # Hidden slides
                try:
                    if slide._element.attrib.get("show", "1") == "0":
                        hidden["hidden_slides"].append({"slide": i + 1})
                        self.findings.append({
                            "type": "HIDDEN_SLIDE",
                            "detail": f"Hidden slide: #{i+1}",
                            "severity": "medium"
                        })
                except:
                    pass

                # Notes
                if slide.has_notes_slide and slide.notes_slide:
                    notes_text = slide.notes_slide.notes_text_frame.text
                    if notes_text.strip():
                        hidden["comments"].append({
                            "slide": i + 1,
                            "type": "speaker_notes",
                            "text": notes_text[:300],
                        })

                # White/invisible text in shapes
                for shape in slide.shapes:
                    if shape.has_text_frame:
                        for para in shape.text_frame.paragraphs:
                            for run in para.runs:
                                if run.font and run.font.color and run.font.color.rgb:
                                    rgb = str(run.font.color.rgb)
                                    if rgb.upper() in ("FFFFFF", "FEFEFE"):
                                        if run.text.strip():
                                            hidden["white_text"].append({
                                                "slide": i + 1,
                                                "text": run.text[:100],
                                            })
                                            self.findings.append({
                                                "type": "WHITE_TEXT",
                                                "detail": f"Invisible text on slide {i+1}: {run.text[:60]}",
                                                "severity": "high"
                                            })
        except Exception as e:
            self.log(f"PPTX hidden content scan: {e}", "WARN")

    def _hidden_ooxml_raw(self, hidden):
        """Scan raw XML for tracked changes, custom XML, etc."""
        try:
            with zipfile.ZipFile(self.path, "r") as z:
                for name in z.namelist():
                    if not name.endswith(".xml") and not name.endswith(".rels"):
                        continue
                    try:
                        xml = z.read(name).decode("utf-8", "replace")

                        # Tracked changes (w:ins, w:del)
                        ins_count = xml.count("<w:ins ")
                        del_count = xml.count("<w:del ")
                        if ins_count + del_count > 0 and "document.xml" in name.lower():
                            root = ET.fromstring(z.read(name))
                            for change_type, tag in [("insertion", f"{{{self.NS['w']}}}ins"),
                                                      ("deletion", f"{{{self.NS['w']}}}del")]:
                                for el in root.findall(f".//{tag}"):
                                    author = el.attrib.get(f"{{{self.NS['w']}}}author", "?")
                                    date = el.attrib.get(f"{{{self.NS['w']}}}date", "?")
                                    texts = []
                                    for t in el.findall(f".//{{{self.NS['w']}}}t"):
                                        if t.text: texts.append(t.text)
                                    if texts:
                                        hidden["tracked_changes"].append({
                                            "type": change_type,
                                            "author": author,
                                            "date": date[:19],
                                            "text": " ".join(texts)[:200],
                                        })
                            if ins_count + del_count > 0:
                                self.findings.append({
                                    "type": "TRACKED_CHANGES",
                                    "detail": f"Tracked changes: {ins_count} insertion(s), {del_count} deletion(s)",
                                    "severity": "medium"
                                })
                    except:
                        pass
        except:
            pass

    # ── Embedded Objects ─────────────────────────────────────────────────────
    def find_embedded_objects(self):
        objects = []

        if self.is_ooxml:
            try:
                with zipfile.ZipFile(self.path, "r") as z:
                    for name in z.namelist():
                        lower = name.lower()
                        obj_type = None
                        if "/embeddings/" in lower:
                            obj_type = "OLE Embedding"
                        elif "/media/" in lower:
                            obj_type = "Media"
                        elif "activex" in lower:
                            obj_type = "ActiveX Control"
                            self.findings.append({
                                "type": "ACTIVEX",
                                "detail": f"ActiveX control: {name}",
                                "severity": "high"
                            })
                        elif lower.endswith((".exe", ".bat", ".cmd", ".ps1", ".vbs", ".js", ".wsf")):
                            obj_type = "Executable"
                            self.findings.append({
                                "type": "EXECUTABLE",
                                "detail": f"Embedded executable: {name}",
                                "severity": "high"
                            })
                        elif "externallinks" in lower:
                            obj_type = "External Link"
                            self.findings.append({
                                "type": "EXTERNAL_LINK",
                                "detail": f"External data link: {name}",
                                "severity": "medium"
                            })

                        if obj_type:
                            info = z.getinfo(name)
                            objects.append({
                                "name": name,
                                "type": obj_type,
                                "size": info.file_size,
                                "compressed": info.compress_size,
                            })

                    # Relationships for external references
                    for name in z.namelist():
                        if name.endswith(".rels"):
                            try:
                                root = ET.fromstring(z.read(name))
                                for rel in root:
                                    target_mode = rel.attrib.get("TargetMode", "")
                                    target = rel.attrib.get("Target", "")
                                    if target_mode == "External":
                                        objects.append({
                                            "name": target[:80],
                                            "type": "External Reference",
                                            "size": 0,
                                            "rel_type": rel.attrib.get("Type", "").split("/")[-1],
                                        })
                                        if target.lower().startswith("http"):
                                            self.findings.append({
                                                "type": "EXTERNAL_URL",
                                                "detail": f"External URL reference: {target[:80]}",
                                                "severity": "medium"
                                            })
                            except:
                                pass
            except:
                pass

        # OLE embedded objects
        if self.is_ole and OLEFILE_OK:
            try:
                ole = olefile.OleFileIO(self.path)
                for stream in ole.listdir():
                    path_str = "/".join(stream)
                    size = ole.get_size(path_str)
                    objects.append({
                        "name": path_str,
                        "type": "OLE Stream",
                        "size": size,
                    })
                ole.close()
            except:
                pass

        return objects

    # ── Revision History Extraction ──────────────────────────────────────────
    def extract_revision_history(self):
        revisions = []

        if not self.is_ooxml:
            return revisions

        try:
            with zipfile.ZipFile(self.path, "r") as z:
                # Check for revision save IDs in core metadata
                for core_path in ["docProps/core.xml"]:
                    if core_path in z.namelist():
                        xml = z.read(core_path)
                        root = ET.fromstring(xml)
                        rev = root.find(f"{{{self.NS['cp']}}}revision")
                        created = root.find(f"{{{self.NS['dcterms']}}}created")
                        modified = root.find(f"{{{self.NS['dcterms']}}}modified")
                        creator = root.find(f"{{{self.NS['dc']}}}creator")
                        last_mod_by = root.find(f"{{{self.NS['cp']}}}lastModifiedBy")

                        if rev is not None and rev.text:
                            revisions.append({
                                "type": "revision_count",
                                "value": rev.text,
                                "detail": f"Document has been saved {rev.text} time(s)"
                            })
                            try:
                                rev_count = int(rev.text)
                                if rev_count > 50:
                                    self.findings.append({
                                        "type": "REVISION_COUNT",
                                        "detail": f"High revision count: {rev_count} saves",
                                        "severity": "low"
                                    })
                            except:
                                pass

                        if created is not None and modified is not None:
                            revisions.append({
                                "type": "timeline",
                                "created": created.text,
                                "modified": modified.text,
                                "creator": creator.text if creator is not None else "?",
                                "last_editor": last_mod_by.text if last_mod_by is not None else "?",
                            })

                # Scan for rsid (revision session IDs) in document.xml
                for doc_path in ["word/document.xml", "xl/workbook.xml", "ppt/presentation.xml"]:
                    if doc_path not in z.namelist():
                        continue
                    xml = z.read(doc_path).decode("utf-8", "replace")
                    rsids = set(re.findall(r'w:rsid[A-Za-z]*="([0-9A-Fa-f]+)"', xml))
                    if rsids:
                        revisions.append({
                            "type": "rsid_sessions",
                            "count": len(rsids),
                            "samples": sorted(rsids)[:20],
                            "detail": f"{len(rsids)} unique editing session(s) detected"
                        })
                        if len(rsids) > 20:
                            self.findings.append({
                                "type": "MANY_SESSIONS",
                                "detail": f"{len(rsids)} unique editing sessions (RSIDs)",
                                "severity": "low"
                            })
                    break
        except Exception as e:
            self.log(f"Revision history extraction: {e}", "WARN")

        return revisions

    # ── Entropy Analysis ─────────────────────────────────────────────────────
    def entropy_analysis(self):
        results = []

        with open(self.path, "rb") as f:
            full_data = f.read()

        results.append({
            "name": "FULL FILE",
            "size": len(full_data),
            "entropy": entropy(full_data),
        })

        # High entropy = possibly encrypted or compressed payload
        file_ent = entropy(full_data)
        if file_ent > 7.5:
            self.findings.append({
                "type": "ENTROPY",
                "detail": f"Very high file entropy ({file_ent:.3f}/8.0) — possible encryption or packing",
                "severity": "medium"
            })

        # Per-ZIP-entry entropy
        if self.is_ooxml:
            try:
                with zipfile.ZipFile(self.path, "r") as z:
                    for info in z.infolist():
                        if info.file_size == 0 or info.is_dir():
                            continue
                        try:
                            data = z.read(info.filename)
                            ent = entropy(data)
                            results.append({
                                "name": info.filename,
                                "size": info.file_size,
                                "entropy": ent,
                            })
                            if ent > 7.8 and info.file_size > 1000:
                                if not info.filename.endswith((".png", ".jpg", ".jpeg", ".gif", ".emf", ".wmf")):
                                    self.findings.append({
                                        "type": "ENTROPY",
                                        "detail": f"High entropy in {info.filename} ({ent:.3f})",
                                        "severity": "medium"
                                    })
                        except:
                            pass
            except:
                pass

        return results

    # ── String Extraction ────────────────────────────────────────────────────
    def extract_strings(self, min_len=5):
        strings = []
        with open(self.path, "rb") as f:
            data = f.read()
        current = ""
        for b in data:
            if 0x20 <= b < 0x7F:
                current += chr(b)
            else:
                if len(current) >= min_len:
                    strings.append(current)
                current = ""
        if len(current) >= min_len:
            strings.append(current)

        # Flag interesting strings
        suspicious_patterns = [
            (r'https?://\S+', "URL"),
            (r'[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Z|a-z]{2,}', "Email"),
            (r'\\\\[A-Za-z0-9_.]+\\', "UNC Path"),
            (r'[A-Z]:\\[^\s]{3,}', "File Path"),
            (r'cmd\.exe|powershell|wscript|cscript|mshta|certutil', "Suspicious Command"),
            (r'password|passwd|secret|token|api.?key', "Credential Keyword"),
        ]
        for s in strings:
            for pattern, label in suspicious_patterns:
                if re.search(pattern, s, re.IGNORECASE):
                    self.findings.append({
                        "type": "STRING",
                        "detail": f"[{label}] {s[:80]}",
                        "severity": "medium" if label in ("Suspicious Command", "Credential Keyword") else "low"
                    })
                    break

        return strings

    # ── Relationship Mapping ─────────────────────────────────────────────────
    def map_relationships(self):
        rels = []
        if not self.is_ooxml:
            return rels
        try:
            with zipfile.ZipFile(self.path, "r") as z:
                for name in z.namelist():
                    if not name.endswith(".rels"):
                        continue
                    try:
                        root = ET.fromstring(z.read(name))
                        for rel in root:
                            rels.append({
                                "source": name,
                                "id": rel.attrib.get("Id", "?"),
                                "type": rel.attrib.get("Type", "?").split("/")[-1],
                                "target": rel.attrib.get("Target", "?"),
                                "mode": rel.attrib.get("TargetMode", "Internal"),
                            })
                    except:
                        pass
        except:
            pass
        return rels

    # ── Full Analysis Pipeline ───────────────────────────────────────────────
    def run_full_analysis(self, progress_fn=None):
        report = {
            "file": self.path,
            "doc_type": self.doc_type,
            "timestamp": datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        }

        if progress_fn: progress_fn("ZIP/OOXML structure...", 5)
        report["zip_structure"] = self.analyze_zip_structure()

        if progress_fn: progress_fn("Extracting metadata...", 15)
        report["metadata"] = self.extract_metadata()

        if progress_fn: progress_fn("Scanning macros/VBA...", 25)
        report["macros"] = self.analyze_macros()

        if progress_fn: progress_fn("Finding hidden content...", 40)
        report["hidden"] = self.find_hidden_content()

        if progress_fn: progress_fn("Embedded objects...", 55)
        report["objects"] = self.find_embedded_objects()

        if progress_fn: progress_fn("Revision history...", 65)
        report["revisions"] = self.extract_revision_history()

        if progress_fn: progress_fn("Entropy analysis...", 78)
        report["entropy"] = self.entropy_analysis()

        if progress_fn: progress_fn("Relationship mapping...", 88)
        report["relationships"] = self.map_relationships()

        report["findings"] = self.findings

        if progress_fn: progress_fn("Analysis complete", 100)
        return report


# ─── Splash Screen ────────────────────────────────────────────────────────────
